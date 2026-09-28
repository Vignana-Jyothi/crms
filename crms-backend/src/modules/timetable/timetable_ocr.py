#!/usr/bin/env python3
"""
timetable_ocr.py - Extract structured data from timetable-style documents.

Supports: images (png/jpg/...), PDFs (digital or scanned), Excel (.xlsx/.xlsm) and CSV.

What it extracts (see the sample: VNR VJIET "Timetable" sheet):
  * header block  -> titles + key/value metadata (Academic Year, Regulation, Room No, w.e.f ...)
  * timetable grid -> {day: [{periods, start, end, subject}]}  (merged cells / labs spanning
                      several periods are handled, LUNCH column is detected)
  * course table  -> list of records (Course Code, Name, Room, Coordinator ...)
  * footer text   -> Coordinator / I-C Timetables / HOD, etc.

Outputs, per input file:  <name>_rows.csv (form-ready rows), <name>.json and <name>.xlsx

Install:
    pip install opencv-python-headless pytesseract numpy openpyxl pdfplumber pypdfium2
    # plus the Tesseract engine itself:
    #   Ubuntu/Debian: sudo apt install tesseract-ocr
    #   macOS:         brew install tesseract
    #   Windows:       https://github.com/UB-Mannheim/tesseract/wiki

Usage:
    python timetable_ocr.py timetable.png
    python timetable_ocr.py a.pdf b.xlsx c.jpg -o out/ --debug
"""
from __future__ import annotations

import argparse
import csv
import difflib
import json
import re
import statistics
import sys
from pathlib import Path

import numpy as np

IMG_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
DAYS = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
TIME_RE = re.compile(r"(\d{1,2})\s*[.:]\s*(\d{2})\s*([AP])\.?\s*M", re.I)

# A "grid cell" is a dict: {r0, r1, c0, c1, text}  (row/col spans, r1/c1 exclusive)


# --------------------------------------------------------------------------- helpers
def clean(t: str) -> str:
    t = re.sub(r"[£€](?=\d)", "E", str(t))  # room numbers like E105 are often read as £105
    return re.sub(r"\s+", " ", t.replace("|", " ")).strip()


def _cluster(vals, tol):
    vals = sorted(vals)
    groups = [[vals[0]]]
    for v in vals[1:]:
        if v - groups[-1][-1] <= tol:
            groups[-1].append(v)
        else:
            groups.append([v])
    return [sum(g) / len(g) for g in groups]


def boxes_to_grid(boxes, tol):
    """boxes: [(x0, y0, x1, y1, text)] in pixels/points -> grid cells with row/col spans."""
    xs = _cluster([b[0] for b in boxes] + [b[2] for b in boxes], tol)
    ys = _cluster([b[1] for b in boxes] + [b[3] for b in boxes], tol)

    def near(edges, v):
        return min(range(len(edges)), key=lambda i: abs(edges[i] - v))

    grid = []
    for x0, y0, x1, y1, t in boxes:
        c0, c1, r0, r1 = near(xs, x0), near(xs, x1), near(ys, y0), near(ys, y1)
        grid.append(dict(r0=r0, r1=max(r1, r0 + 1), c0=c0, c1=max(c1, c0 + 1), text=t))
    return grid


def words_to_lines(words, gap_factor=1.5):
    """words: [{text,x0,x1,y0,y1}] -> list of lines; each line = list of text segments.
    Words separated by a wide horizontal gap become separate segments (label / value pairs)."""
    words = sorted(words, key=lambda w: (w["y0"] + w["y1"]) / 2)
    lines = []
    for w in words:
        yc, h = (w["y0"] + w["y1"]) / 2, w["y1"] - w["y0"]
        if lines and abs(yc - lines[-1]["yc"]) < 0.6 * max(h, lines[-1]["h"]):
            ln = lines[-1]
            ln["w"].append(w)
            ln["yc"] = (ln["yc"] * (len(ln["w"]) - 1) + yc) / len(ln["w"])
            ln["h"] = max(ln["h"], h)
        else:
            lines.append(dict(w=[w], yc=yc, h=h))
    out = []
    for ln in lines:
        ws = sorted(ln["w"], key=lambda w: w["x0"])
        gap = gap_factor * statistics.median(w["y1"] - w["y0"] for w in ws)
        segs, cur = [], [ws[0]]
        for a, b in zip(ws, ws[1:]):
            if b["x0"] - a["x1"] > gap:
                segs.append(cur)
                cur = [b]
            else:
                cur.append(b)
        segs.append(cur)
        out.append([clean(" ".join(w["text"] for w in s)) for s in segs])
    return out


_LABELS = ["academic year", "regulation", "section", "programme", "class", "room no", "semester", "branch",
           "w e f", "department", "year"]


def _is_label(s):
    t = re.sub(r"[^a-z ]", " ", s.lower()).split()
    t = " ".join(t)
    return bool(t) and any(difflib.SequenceMatcher(None, t, l).ratio() >= 0.85 for l in _LABELS)


def parse_lines(lines):
    """Turn text lines into (metadata dict, titles, other text)."""
    meta, titles, other = {}, [], []
    for segs in lines:
        i = 0
        while i < len(segs):
            s = segs[i]
            m = re.match(r"^([^:]{2,40}?)\s*:\s*(.*)$", s)
            if m:
                k, v = m.group(1).strip(), m.group(2).strip()
                if not v and i + 1 < len(segs) and ":" not in segs[i + 1]:
                    v = segs[i + 1]
                    i += 1
                meta[k] = v
            elif _is_label(s) and i + 1 < len(segs):  # label whose ':' was lost by OCR
                meta[re.sub(r"[.:;,\s]+$", "", s)] = segs[i + 1]
                i += 1
            elif s:
                (titles if len(segs) == 1 else other).append(s)
            i += 1
    return meta, titles, other


# --------------------------------------------------------------------------- parsers
def _fmt_time(m):
    return f"{int(m[0])}:{m[1]} {m[2].upper()}M"


def _infer_missing_periods(hdr):
    """OCR often drops a lone period digit. An un-numbered column that sits in a numbering gap
    (e.g. between 4 and 6) is period 5; otherwise it is a break (lunch). Also fills missing end times."""
    i = 0
    while i < len(hdr):
        if hdr[i]["period"] is None:
            j = i
            while j < len(hdr) and hdr[j]["period"] is None:
                j += 1
            prev = hdr[i - 1]["period"] if i > 0 else 0
            nxt = hdr[j]["period"] if j < len(hdr) else None
            n = j - i
            if nxt is not None and nxt - prev - 1 >= n:
                for k in range(i, j):
                    hdr[k]["period"], hdr[k]["is_break"] = prev + 1 + (k - i), False
        i = max(i + 1, 0) if hdr[i]["period"] is not None else j
    for a, b in zip(hdr, hdr[1:]):
        if not a["end"]:
            a["end"] = b["start"]


def parse_timetable(cells):
    """Find a Day x Period grid. Returns dict or None."""
    hdr = []
    for c in cells:
        times = TIME_RE.findall(c["text"])
        if times:
            rest = TIME_RE.sub(" ", c["text"])
            p = re.findall(r"\b(\d{1,2})\b", rest)
            hdr.append(dict(c0=c["c0"], c1=c["c1"], r0=c["r0"], r1=c["r1"],
                            period=int(p[-1]) if p else None,
                            start=_fmt_time(times[0]), end=_fmt_time(times[-1]) if len(times) > 1 else "",
                            is_break=not p))
    if len(hdr) < 3:
        return None
    hdr_r0 = min(h["r0"] for h in hdr)
    hdr = sorted([h for h in hdr if h["r0"] == hdr_r0], key=lambda h: h["c0"])
    _infer_missing_periods(hdr)
    min_c0 = min(h["c0"] for h in hdr)

    day_cells = []
    for c in cells:
        if c["c1"] <= min_c0 and c["r0"] > hdr_r0:
            m = difflib.get_close_matches(c["text"].strip().lower(), DAYS, n=1, cutoff=0.6)
            if m:
                day_cells.append((m[0].capitalize(), c))
    if len(day_cells) < 2:
        return None

    days = {}
    for name, d in sorted(day_cells, key=lambda x: x[1]["r0"]):
        entries = []
        for c in sorted(cells, key=lambda c: c["c0"]):
            if c["c0"] < min_c0 or not (c["r0"] < d["r1"] and c["r1"] > d["r0"]):
                continue
            cols = [h for h in hdr if h["c0"] < c["c1"] and h["c1"] > c["c0"]]
            per = [h for h in cols if not h["is_break"]]
            if not per or not c["text"].strip():
                continue
            entries.append(dict(periods=[h["period"] for h in per], start=per[0]["start"],
                                end=per[-1]["end"], subject=c["text"].strip(),
                                _c0=per[0]["c0"], _c1=per[-1]["c1"]))
        days[name] = entries

    return dict(
        columns=[dict(period=h["period"], start=h["start"], end=h["end"], is_break=h["is_break"],
                      _c0=h["c0"], _c1=h["c1"]) for h in hdr],
        days=days, _header_r0=hdr_r0, _last_r1=max(d["r1"] for _, d in day_cells))


def _is_coordinator_header(text):
    """Header OCR is often garbled, so match fuzzily against the expected wording."""
    t = re.sub(r"[^a-z ]", "", text.lower())
    if re.search(r"coordinator|faculty", t):
        return True
    sim = lambda ref: difflib.SequenceMatcher(None, t, ref).ratio()
    return sim("name of the course coordinator faculty") > max(0.5, sim("name of the course short name") + 0.05)


def parse_courses(cells):
    """Find a 'Course Code ...' table. Columns come from the *body* cells (headers can be garbled or
    partly hidden); header text is only used for names. Side-by-side tables are split after the
    faculty column. Returns dict(tables=[[record, ...], ...]) or None."""
    h0 = next((c for c in cells if re.search(r"course\s*code", c["text"], re.I)), None)
    if not h0:
        return None
    hdr = sorted([c for c in cells if c["r0"] == h0["r0"]], key=lambda c: c["c0"])
    body = [c for c in cells if c["r0"] >= h0["r1"] and c["text"].strip()]
    if not body:
        return None
    starts = sorted({c["c0"] for c in body})
    nxt = lambda i: starts[i + 1] if i + 1 < len(starts) else 10 ** 6

    def head(i):
        best, bo = None, 0
        for h in hdr:
            o = min(h["c1"], nxt(i)) - max(h["c0"], starts[i])
            if o > bo:
                best, bo = h, o
        return best["text"] if best else ""

    title = re.compile(r"(?i)^(mr|mrs|ms|dr|prof)\b")

    def faculty_col(i):
        vals = [c for c in body if c["c0"] == starts[i]]
        return len(vals) >= 3 and sum(bool(title.match(c["text"])) for c in vals) / len(vals) >= 0.5

    names, seen, group_of, g = [], {}, {}, 0
    for i, st in enumerate(starts):
        n = head(i) or f"col_{st}"
        seen[n] = seen.get(n, 0) + 1
        names.append(n if seen[n] == 1 else f"{n} ({seen[n]})")
        group_of[i] = g
        if faculty_col(i) or _is_coordinator_header(head(i)):
            g += 1
    tables = [[] for _ in range(max(group_of.values()) + 1)]
    for r in sorted({c["r0"] for c in body}):
        recs = [dict() for _ in tables]
        for c in body:
            if c["r0"] == r:
                i = starts.index(c["c0"])
                recs[group_of[i]][names[i]] = c["text"].strip()
        for gi, rec in enumerate(recs):
            if rec:
                tables[gi].append(rec)
    return dict(tables=[t for t in tables if t], _header_r0=h0["r0"])


def snap_to_known(text, known, cutoff=0.8):
    parts = [p.strip() for p in re.split(r"\s*/\s*", text) if p.strip()]
    out = []
    for p in parts:
        u = p.upper()
        if u in known or _norm(u) in {_norm(k) for k in known}:
            out.append(p)
        elif len(u) <= 2:  # short codes: accept a unique candidate that differs by one character
            m = [k for k in known if len(k) == len(u) and sum(a != b for a, b in zip(k, u)) <= 1]
            out.append(m[0] if len(m) == 1 else p)
        else:
            m = difflib.get_close_matches(u, sorted(known), n=1, cutoff=max(cutoff, 0.86) if len(u) <= 4 else cutoff)
            out.append(m[0] if m else p)
    return out


def cells_to_text_lines(cells):
    rows = {}
    for c in cells:
        rows.setdefault(c["r0"], []).append(c)
    return [[clean(c["text"]) for c in sorted(v, key=lambda c: c["c0"]) if c["text"].strip()]
            for _, v in sorted(rows.items())]


def cells_to_matrix(cells):
    if not cells:
        return []
    nr, nc = max(c["r1"] for c in cells), max(c["c1"] for c in cells)
    m = [[""] * nc for _ in range(nr)]
    for c in cells:
        for r in range(c["r0"], c["r1"]):
            for k in range(c["c0"], c["c1"]):
                m[r][k] = c["text"]
    return [row for row in m if any(row)]


def process_page(source, page, blocks, text_lines, excel_like=False):
    res = dict(source=source, page=page, titles=[], metadata={}, other_text=[],
               timetables=[], course_tables=[], other_tables=[])
    lines = list(text_lines)
    queue = list(blocks)
    known = set()
    while queue:
        blk = queue.pop(0)
        tt = parse_timetable(blk)
        ct = None if tt else parse_courses(blk)
        found = tt or ct
        if found:
            hr = found["_header_r0"]
            lines += cells_to_text_lines([c for c in blk if c["r1"] <= hr])
            if tt:
                res["timetables"].append(tt)
                queue.append([c for c in blk if c["r0"] >= tt["_last_r1"]])
            else:
                res["course_tables"].append(ct)
        elif blk:
            few = max(len(r) for r in cells_to_text_lines(blk)) <= 4 and len({c["r0"] for c in blk}) <= 8
            if excel_like and few:
                lines += cells_to_text_lines(blk)
            else:
                res["other_tables"].append(cells_to_matrix(blk))

    # vocabulary for OCR clean-up: short names in "(...)" + category / activity names
    for ct in res["course_tables"]:
        for gi, tbl in enumerate(ct["tables"]):
            for rec in tbl:
                for k, v in rec.items():
                    if gi == 0:
                        m = re.search(r"\(([A-Za-z0-9 &/\-]{2,25})\)\s*$", v)
                        if m:
                            known.add(m.group(1).strip().upper())
                        if re.search(r"course", k, re.I) and "code" not in k.lower():
                            words = [w for w in re.sub(r"\(.*?\)", "", v).split()
                                     if w.lower() not in {"and", "of", "the", "for", "&"}]
                            if 2 <= len(words) <= 4:  # acronym, e.g. Design Thinking -> DT
                                known.add("".join(w[0] for w in words).upper())
                    elif re.search(r"categor|activity|course", k, re.I):
                        known.update(p.strip().upper() for p in re.split(r"/", v) if p.strip())
    for tt in res["timetables"]:
        for entries in tt["days"].values():
            for e in entries:
                raw = re.sub(r"\s*&\s*", " & ", e["subject"]).strip()  # 'MR&MM' -> 'MR & MM'
                opts = snap_to_known(raw, known) if known else [p.strip() for p in raw.split("/")]
                opts = [o.upper() if len(o) <= 4 and o.isalpha() and o.upper() not in known and known else o
                        for o in opts]
                e["options"] = opts
                e["subject"] = " / ".join(opts)
                if e["subject"].replace(" ", "") != raw.replace(" ", ""):
                    e["raw_ocr"] = raw

    # one unmatched "X LAB" token + exactly one never-used known lab => OCR confusion (e.g. SDM -> USE)
    used = {o.upper() for tt in res["timetables"] for es in tt["days"].values() for e in es for o in e["options"]}
    unused = [k for k in known if k.endswith(" LAB") and k not in used]
    unknown = {o for o in used if o.endswith(" LAB") and o not in known}
    if len(unused) == 1 and len(unknown) == 1:
        bad, good = next(iter(unknown)), unused[0]
        for tt in res["timetables"]:
            for es in tt["days"].values():
                for e in es:
                    if any(o.upper() == bad for o in e["options"]):
                        e.setdefault("raw_ocr", e["subject"])
                        e["options"] = [good if o.upper() == bad else o for o in e["options"]]
                        e["subject"] = " / ".join(e["options"])
    res["_activities"] = sorted({p.strip().upper() for ct in res["course_tables"] for tbl in ct["tables"][1:]
                                 for rec in tbl for v in rec.values() for p in v.split("/") if p.strip()})

    meta, titles, other = parse_lines(lines)
    res.update(metadata=meta, titles=titles, other_text=other)
    return res


# --------------------------------------------------------------------------- image / OCR
def _check_tesseract():
    import pytesseract
    try:
        pytesseract.get_tesseract_version()
    except Exception:
        sys.exit("Tesseract engine not found. Install it (apt install tesseract-ocr / brew install "
                 "tesseract / Windows installer) and make sure it is on PATH.")
    return pytesseract


def _ocr_cell(gray, box, lang, psm):
    """OCR one cell. Retries with larger scale / binarisation when confidence is low."""
    import cv2
    pytesseract = _check_tesseract()
    x0, y0, x1, y1 = box
    pad = 3
    crop = gray[y0 + pad:y1 - pad, x0 + pad:x1 - pad]
    if crop.size == 0 or (crop < 140).mean() < 0.004:
        return ""

    def run(img, p):
        img = cv2.copyMakeBorder(img, 12, 12, 12, 12, cv2.BORDER_CONSTANT, value=255)
        d = pytesseract.image_to_data(img, lang=lang, config=f"--psm {p}", output_type=pytesseract.Output.DICT)
        idx = [i for i, t in enumerate(d["text"]) if t.strip() and float(d["conf"][i]) >= 0]
        if not idx:
            return "", 0.0
        conf = sum(float(d["conf"][i]) for i in idx) / len(idx)
        rows = {}
        for i in idx:
            rows.setdefault((d["block_num"][i], d["par_num"][i], d["line_num"][i]), []).append(d["text"][i])
        lines = [" ".join(v) for v in rows.values()]
        if len(lines) >= 3 and all(len(l) <= 1 for l in lines):  # vertical text: L U N C H
            return "".join(lines), conf
        return clean(" ".join(lines)), conf

    def scaled(target):
        f = min(4.0, target / crop.shape[0])
        return cv2.resize(crop, None, fx=f, fy=f, interpolation=cv2.INTER_CUBIC) if f > 1 else crop

    attempts = [(scaled(90), psm)]
    best_t, best_c = "", -1.0
    for n, (img, p) in enumerate(attempts):
        t, c = run(img, p)
        if c > best_c:
            best_t, best_c = t, c
    if best_c < 80:  # low confidence -> try bigger, sharper, alternative segmentation modes
        for target in (150, 220):
            big = scaled(target)
            sharp = cv2.threshold(cv2.GaussianBlur(big, (3, 3), 0), 0, 255, cv2.THRESH_BINARY | cv2.THRESH_OTSU)[1]
            for img, p in ((big, psm), (sharp, psm), (sharp, 7)):
                t, c = run(img, p)
                if c > best_c:
                    best_t, best_c = t, c
            if best_c >= 85:
                break
    return best_t


def ocr_image(gray, lang="eng", psm=6, debug_path=None):
    """Returns (blocks, text_lines) for one page image (grayscale ndarray)."""
    import cv2
    pytesseract = _check_tesseract()

    H, W = gray.shape
    if W < 2000:  # small screenshots OCR badly - upscale
        f = min(4.0, 2000 / W)
        gray = cv2.resize(gray, None, fx=f, fy=f, interpolation=cv2.INTER_CUBIC)
        H, W = gray.shape

    # 1. ruling lines -> mask
    bw = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)[1]
    hor = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (max(15, W // 60), 1)))
    ver = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(15, W // 80))))
    lines = cv2.dilate(cv2.bitwise_or(hor, ver), np.ones((3, 3), np.uint8))

    # 2. table regions = large connected line components
    n, _, stats, _ = cv2.connectedComponentsWithStats(lines, connectivity=8)
    tables = [tuple(stats[i][:4]) for i in range(1, n) if stats[i][2] > 0.15 * W and stats[i][3] > 0.04 * H]
    tables.sort(key=lambda t: t[1])

    # 3. cells = enclosed white regions
    # RETR_CCOMP: cells sit inside the "hole" of the page background, so they show up as
    # top-level contours (parent == -1); the holes themselves (table outlines) are skipped.
    cnts, hier = cv2.findContours(cv2.bitwise_not(lines), cv2.RETR_CCOMP, cv2.CHAIN_APPROX_SIMPLE)
    cell_boxes = []
    for c, hh in zip(cnts, hier[0]):
        if hh[3] != -1:
            continue
        x, y, w, h = cv2.boundingRect(c)
        if w < 0.012 * W or h < 0.008 * H or w > 0.95 * W or h > 0.6 * H:
            continue
        if cv2.contourArea(c) / (w * h) < 0.7:
            continue
        cell_boxes.append((x, y, x + w, y + h))

    blocks, dbg = [], cv2.cvtColor(gray, cv2.COLOR_GRAY2BGR)
    for tx, ty, tw, th in tables:
        inside = [b for b in cell_boxes
                  if tx - 5 <= (b[0] + b[2]) / 2 <= tx + tw + 5 and ty - 5 <= (b[1] + b[3]) / 2 <= ty + th + 5]
        if not inside:
            continue
        boxes = [(*b, _ocr_cell(gray, b, lang, psm)) for b in inside]
        blocks.append(boxes_to_grid(boxes, tol=max(4, W // 200)))
        for b in inside:
            cv2.rectangle(dbg, b[:2], b[2:], (0, 0, 255), 1)
        cv2.rectangle(dbg, (tx, ty), (tx + tw, ty + th), (0, 160, 0), 2)
    if debug_path:
        cv2.imwrite(str(debug_path), dbg)

    # 4. text outside tables (titles, metadata, footer)
    masked = gray.copy()
    for tx, ty, tw, th in tables:
        cv2.rectangle(masked, (tx - 4, ty - 4), (tx + tw + 4, ty + th + 4), 255, -1)
    masked = cv2.resize(masked, None, fx=1.6, fy=1.6, interpolation=cv2.INTER_CUBIC)  # coordinates only used relatively
    d = pytesseract.image_to_data(masked, lang=lang, config="--psm 11", output_type=pytesseract.Output.DICT)
    words = [dict(text=d["text"][i], x0=d["left"][i], x1=d["left"][i] + d["width"][i],
                  y0=d["top"][i], y1=d["top"][i] + d["height"][i])
             for i in range(len(d["text"])) if d["text"][i].strip() and float(d["conf"][i]) > 20]
    return blocks, (words_to_lines(words) if words else [])


# --------------------------------------------------------------------------- readers
def read_image(path, opts):
    import cv2
    img = cv2.imread(str(path), cv2.IMREAD_GRAYSCALE)
    if img is None:
        sys.exit(f"Cannot read image: {path}")
    dbg = Path(opts.out) / f"{path.stem}_debug.png" if opts.debug else None
    blocks, lines = ocr_image(img, opts.lang, opts.psm, dbg)
    return [process_page(path.name, 1, blocks, lines)]


def read_pdf(path, opts):
    import pdfplumber
    results = []
    with pdfplumber.open(str(path)) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            blocks, lines = [], []
            if len((page.extract_text() or "").strip()) > 20:  # digital PDF: use the text layer
                tables = page.find_tables()
                bboxes = [t.bbox for t in tables]
                for t in tables:
                    boxes = []
                    for x0, y0, x1, y1 in t.cells:
                        bb = (max(x0, page.bbox[0]), max(y0, page.bbox[1]), min(x1, page.bbox[2]), min(y1, page.bbox[3]))
                        txt = clean(page.within_bbox(bb).extract_text() or "")
                        boxes.append((*bb, txt))
                    if boxes:
                        blocks.append(boxes_to_grid(boxes, tol=2.0))

                def outside(o):
                    if o.get("object_type") != "char":
                        return True
                    cx, cy = (o["x0"] + o["x1"]) / 2, (o["top"] + o["bottom"]) / 2
                    return not any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in bboxes)

                ws = page.filter(outside).extract_words(x_tolerance=1.5)
                words = [dict(text=w["text"], x0=w["x0"], x1=w["x1"], y0=w["top"], y1=w["bottom"]) for w in ws]
                lines = words_to_lines(words) if words else []
            if not blocks:  # scanned PDF (or no ruling lines): rasterise and OCR
                import pypdfium2 as pdfium
                doc = pdfium.PdfDocument(str(path))
                img = np.array(doc[pno - 1].render(scale=300 / 72).to_pil().convert("L"))
                dbg = Path(opts.out) / f"{path.stem}_p{pno}_debug.png" if opts.debug else None
                blocks, lines = ocr_image(img, opts.lang, opts.psm, dbg)
            results.append(process_page(path.name, pno, blocks, lines))
    return results


def _rows_to_cells(rows, merged=None):
    cells = []
    covered, anchors = set(), {}
    for rng in merged or []:
        anchors[(rng.min_row, rng.min_col)] = rng
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                covered.add((r, c))
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            if (r, c) in anchors:
                g = anchors[(r, c)]
                r1, c1 = g.max_row + 1, g.max_col + 1
            elif (r, c) in covered or v is None or str(v).strip() == "":
                continue
            else:
                r1, c1 = r + 1, c + 1
            cells.append(dict(r0=r, r1=r1, c0=c, c1=c1, text=clean(v if v is not None else "")))
    return cells


def _split_blocks(cells):
    """Split a sheet into blocks separated by fully blank rows."""
    occ = sorted({r for c in cells for r in range(c["r0"], c["r1"])})
    blocks, cur, prev = None, set(), None
    blocks = []
    for r in occ:
        if prev is not None and r > prev + 1:
            blocks.append(cur)
            cur = set()
        cur.add(r)
        prev = r
    if cur:
        blocks.append(cur)
    return [[c for c in cells if c["r0"] in rs] for rs in blocks]


def read_excel(path, opts):
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out = []
    for i, ws in enumerate(wb.worksheets, 1):
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        cells = _rows_to_cells(rows, ws.merged_cells.ranges)
        if cells:
            r = process_page(f"{path.name}:{ws.title}", i, _split_blocks(cells), [], excel_like=True)
            out.append(r)
    return out


def read_csv(path, opts):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    cells = _rows_to_cells(rows)
    return [process_page(path.name, 1, _split_blocks(cells), [], excel_like=True)]


def extract(path: Path, opts):
    ext = path.suffix.lower()
    if ext in IMG_EXT:
        return read_image(path, opts)
    if ext == ".pdf":
        return read_pdf(path, opts)
    if ext in (".xlsx", ".xlsm"):
        return read_excel(path, opts)
    if ext == ".csv":
        return read_csv(path, opts)
    sys.exit(f"Unsupported file type: {ext}  (for .xls / .doc convert to .xlsx / .pdf first)")



# --------------------------------------------------------------------------- upload-form rows
# Output columns match the scheduling form: Day | Start Time | End Time | Subject | Year | Dept |
# Section | Faculty | Classroom
ROW_COLS = ["Day", "Start Time", "End Time", "Subject", "Year", "Dept", "Section", "Faculty", "Classroom"]
_ROMAN = {1: "I", 2: "II", 3: "III", 4: "IV"}


def _norm(s):
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def _mins(t):
    m = re.match(r"(\d{1,2}):(\d{2})\s*([AP])M", t or "", re.I)
    return None if not m else (int(m[1]) % 12 + (12 if m[3].upper() == "P" else 0)) * 60 + int(m[2])


def _hhmm(m):
    return f"{m // 60 % 24:02d}:{m % 60:02d}"


def _year_from_text(text):
    """Roman numeral (with typical OCR confusions) -> int, or None."""
    t = re.sub(r"(?i)year|yr|semester|sem|[^a-z0-9|!]", "", text or "")
    if not t:
        return None
    if t.lower() in ("m", "rn", "iii", "3"):
        return 3
    if t.lower() in ("iv", "4"):
        return 4
    if re.fullmatch(r"[IiLl1|!tT]+", t):
        return min(len(t), 4)
    return None


def _detect_year_dept_section(meta, opts, warn):
    get = lambda pat: next((v for k, v in meta.items() if re.search(pat, k, re.I)), "")
    cls_y = _year_from_text(get(r"^(class|year)"))
    sem = _year_from_text(get(r"sem"))
    sem_y = (sem + 1) // 2 if sem else None
    # roman numerals are the least reliable OCR text: trust "II/III/IV" over a lone stroke
    year = opts.year or (cls_y if cls_y and cls_y >= 2 else sem_y or cls_y)
    if opts.year is None and not (cls_y and cls_y >= 2) and cls_y and sem_y and cls_y != sem_y:
        warn.append(f"Year: class reads as year {cls_y} but semester implies {sem_y}; used {year}. Verify or use --year.")
    if not year:
        warn.append("Year not detected (use --year).")
    branch = get(r"branch|programme")
    words = [w for w in get(r"branch").split() if w.lower() not in {"and", "of", "the", "&"}]
    dept = opts.dept or ("".join(w[0] for w in words).upper() if len(words) > 1 else branch)
    sec = (opts.section or get(r"section")).strip()
    if sec in ("", "-", "--", "---", "\u2014", "\u2013", "NA", "N/A"):
        sec = "A"  # a blank / "--" section means section A
    rm = next((v for k, v in meta.items() if difflib.SequenceMatcher(None, k.lower(), "room no").ratio() > 0.6), "")
    room = opts.room or rm
    yr = "" if not year else (_ROMAN[year] if opts.year_fmt == "roman" else str(year))
    return yr, dept, sec, room


def _classify_record(rec):
    info, rest = {}, []
    for k, v in rec.items():
        if re.match(r"(?i)(mr|mrs|ms|dr|prof)\b", v):
            info["faculty"] = v
        elif re.fullmatch(r"\d{2}[A-Z0-9]{5,8}", v.replace(" ", ""), re.I):
            info["code"] = v
        elif "room" in k.lower():
            info["room"] = v
        else:
            rest.append(v)
    if rest:
        withshort = [v for v in rest if re.search(r"\([^)]{2,25}\)\s*$", v)]
        info["name"] = withshort[0] if withshort else max(rest, key=len)
    return info


def _course_index(res):
    idx = {}
    for ct in res["course_tables"]:
        for rec in (ct["tables"][0] if ct["tables"] else []):
            info = _classify_record(rec)
            name = info.get("name", "")
            keys = set()
            m = re.search(r"\(([^)]{2,25})\)\s*$", name)
            if m:
                keys.add(_norm(m.group(1)))
            words = [w for w in re.sub(r"\(.*?\)", "", name).split() if w.lower() not in {"and", "of", "the", "for", "&"}]
            if 2 <= len(words) <= 4:
                keys.add(_norm("".join(w[0] for w in words)))
            for k in keys:
                idx.setdefault(k, info)
    return idx


def _room_for_day(room_str, day, default=""):
    """'E131(wed)/ E102(thus)' + 'Wednesday' -> 'E131'.  Untagged rooms apply to every day."""
    if not room_str:
        return default
    fix = str.maketrans("iIlLoOsS", "11110055")
    found = re.findall(r"([A-Z][0-9]{2}[0-9iIlLoOsS])\s*(?:[\(\{\[]\s*([A-Za-z]+)\s*[\)\}\]\]])?", room_str)
    tagged = [(r.translate(fix), t.lower()[:3]) for r, t in found if t]
    untagged = [r.translate(fix) for r, t in found if not t]
    for r, t in tagged:
        if t == day.lower()[:3]:
            return r
    if tagged:
        return ""  # course has day-specific rooms and none is for this day
    return untagged[0] if untagged else default


def build_rows(res, opts, warn):
    year, dept, sec, default_room = _detect_year_dept_section(res["metadata"], opts, warn)
    idx = _course_index(res)
    acts = set(res.get("_activities", []))
    rows = []
    for tt in res["timetables"]:
        per_time = {c["period"]: (c["start"], c["end"]) for c in tt["columns"] if not c["is_break"]}
        for day, entries in tt["days"].items():
            for e in entries:
                chunks = [(e["periods"], e["start"], e["end"])]
                if opts.per_period and len(e["periods"]) > 1:
                    chunks = [([p], *per_time.get(p, (e["start"], e["end"]))) for p in e["periods"]]
                for _, st, en in chunks:
                    s0, e0 = _mins(st), _mins(en)
                    if s0 is not None and (e0 is None or e0 <= s0):
                        e0 = (s0 + 60) if e0 is None else e0 + 720  # "12:00 AM" typo -> 12:00 (noon)
                    facs, rooms, subj = [], [], []
                    for o in e["options"]:
                        info = idx.get(_norm(o))
                        subj.append(o)
                        if info:
                            facs.append(info.get("faculty", ""))
                            rooms.append(_room_for_day(info.get("room", ""), day, default_room))
                        elif o.upper() in acts or _norm(o) in {_norm(a) for a in acts}:
                            facs.append("")
                            rooms.append("")  # sports / library / ECA / CCA ... have no room
                        else:
                            facs.append("")
                            rooms.append("")
                            warn.append(f"{day} {st}: no course match for '{o}' (faculty/room left blank)")
                    variants = ([(o, f, r) for o, f, r in zip(subj, facs, rooms)] if opts.split_options
                                else [(" / ".join(subj), " / ".join(f for f in facs if f),
                                       " / ".join(r for r in rooms if r))])
                    for sb, fc, rm in variants:
                        rows.append({"Day": day, "Start Time": _hhmm(s0) if s0 is not None else st,
                                     "End Time": _hhmm(e0) if e0 is not None else en, "Subject": sb,
                                     "Year": year, "Dept": dept, "Section": sec, "Faculty": fc, "Classroom": rm})
    return rows


# --------------------------------------------------------------------------- writers
def _strip_private(o):
    if isinstance(o, dict):
        return {k: _strip_private(v) for k, v in o.items() if not k.startswith("_")}
    if isinstance(o, list):
        return [_strip_private(v) for v in o]
    return o


def write_xlsx(results, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

    wb = Workbook()
    wb.remove(wb.active)
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDE7F5")
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")
    multi = len(results) > 1

    def sheet(name):
        base = re.sub(r"[\[\]\:\*\?\/\\]", "_", name)[:31]
        n, i = base, 2
        while n in wb.sheetnames:
            n = f"{base[:28]}_{i}"
            i += 1
        return wb.create_sheet(n)

    for r in results:
        tag = f"p{r['page']}_" if multi else ""
        if r.get("rows"):
            ws = sheet(f"{tag}Rows")
            ws.append(ROW_COLS)
            for row in r["rows"]:
                ws.append([row[c] for c in ROW_COLS])
            for c in ws[1]:
                c.font, c.fill = bold, fill
            for j, w in enumerate([12, 12, 12, 26, 8, 10, 10, 44, 20], 1):
                ws.column_dimensions[ws.cell(1, j).column_letter].width = w
        ws = sheet(f"{tag}Metadata")
        ws.append(["Field", "Value"])
        for t in r["titles"]:
            ws.append(["Title", t])
        for k, v in r["metadata"].items():
            ws.append([k, v])
        for t in r["other_text"]:
            ws.append(["Other", t])
        for c in ws[1]:
            c.font, c.fill = bold, fill
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 70

        for ti, tt in enumerate(r["timetables"], 1):
            ws = sheet(f"{tag}Timetable" + (f"_{ti}" if ti > 1 else ""))
            cols = tt["columns"]
            for j, c in enumerate(cols, 2):
                lab = "LUNCH" if c["is_break"] else f"Period {c['period']}"
                ws.cell(1, j, f"{lab}\n{c['start']} - {c['end']}".strip(" -"))
            ws.cell(1, 1, "Day")
            for c in ws[1]:
                c.font, c.fill, c.alignment, c.border = bold, fill, wrap, box
            col_of = {c["_c0"]: j for j, c in enumerate(cols, 2)}
            brk = [j for j, c in enumerate(cols, 2) if c["is_break"]]
            for ri, (day, entries) in enumerate(tt["days"].items(), 2):
                ws.cell(ri, 1, day).font = bold
                for e in entries:
                    j0 = col_of.get(e["_c0"], 2)
                    last = [j for j, c in enumerate(cols, 2) if c["_c1"] == e["_c1"]]
                    j1 = last[0] if last else j0
                    ws.cell(ri, j0, e["subject"])
                    if j1 > j0:
                        ws.merge_cells(start_row=ri, start_column=j0, end_row=ri, end_column=j1)
                for j in range(1, len(cols) + 2):
                    ws.cell(ri, j).alignment, ws.cell(ri, j).border = wrap, box
            for j in brk:
                ws.cell(2, j, "L U N C H")
                if len(tt["days"]) > 1:
                    ws.merge_cells(start_row=2, start_column=j, end_row=len(tt["days"]) + 1, end_column=j)
            for j in range(1, len(cols) + 2):
                ws.column_dimensions[ws.cell(1, j).column_letter].width = 18
            ws.row_dimensions[1].height = 32

        for ci, ct in enumerate(r["course_tables"], 1):
            for gi, tbl in enumerate(ct["tables"], 1):
                ws = sheet(f"{tag}Courses" + (f"_{gi}" if gi > 1 else ""))
                heads = list(dict.fromkeys(k for rec in tbl for k in rec))
                ws.append(heads)
                for rec in tbl:
                    ws.append([rec.get(h, "") for h in heads])
                for c in ws[1]:
                    c.font, c.fill = bold, fill
                for j in range(1, len(heads) + 1):
                    ws.column_dimensions[ws.cell(1, j).column_letter].width = 34

        for oi, m in enumerate(r["other_tables"], 1):
            ws = sheet(f"{tag}Raw_{oi}")
            for row in m:
                ws.append(row)
    if not wb.sheetnames:
        wb.create_sheet("Empty")
    wb.save(path)


# --------------------------------------------------------------------------- CLI
def main():
    ap = argparse.ArgumentParser(description="Extract timetables / tables from images, PDFs, Excel, CSV.")
    ap.add_argument("inputs", nargs="+", help="input files")
    ap.add_argument("-o", "--out", default="ocr_output", help="output directory (default: ocr_output)")
    ap.add_argument("--lang", default="eng", help="tesseract language(s), e.g. eng+hin")
    ap.add_argument("--psm", type=int, default=6, help="tesseract page-segmentation mode for cells")
    ap.add_argument("--debug", action="store_true", help="save an image showing detected tables/cells")
    ap.add_argument("--year", type=int, help="force year of study (1-4) instead of auto-detect")
    ap.add_argument("--dept", help="force department code, e.g. CSBS (default: initials of 'Branch')")
    ap.add_argument("--section", help="force section (default: header value; blank or '--' => A)")
    ap.add_argument("--room", help="default classroom when a course has none (default: header 'Room No')")
    ap.add_argument("--year-fmt", choices=["num", "roman"], default="num", help="Year column format: 2 or II")
    ap.add_argument("--split-options", action="store_true", help="'A / B' cells -> one row per option")
    ap.add_argument("--per-period", action="store_true", help="lab blocks spanning periods -> one row per period")
    ap.add_argument("--print", action="store_true", help="print the JSON result to stdout")
    opts = ap.parse_args()

    Path(opts.out).mkdir(parents=True, exist_ok=True)
    for f in opts.inputs:
        p = Path(f)
        if not p.exists():
            print(f"[skip] not found: {p}")
            continue
        print(f"[..] {p.name}")
        results = extract(p, opts)
        warnings = []
        for r in results:
            r["rows"] = build_rows(r, opts, warnings)
        write_xlsx(results, Path(opts.out) / f"{p.stem}.xlsx")
        clean_res = _strip_private(results)
        with open(Path(opts.out) / f"{p.stem}_rows.csv", "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=ROW_COLS)
            w.writeheader()
            for r in results:
                w.writerows(r["rows"])
        (Path(opts.out) / f"{p.stem}.json").write_text(
            json.dumps(clean_res, indent=2, ensure_ascii=False), encoding="utf-8")
        for r in clean_res:
            print(f"     page {r['page']}: {len(r['timetables'])} timetable(s), "
                  f"{sum(len(t) for c in r['course_tables'] for t in c['tables'])} course rows, "
                  f"{len(r['metadata'])} metadata fields, {len(r.get('rows', []))} upload rows")
        for w in dict.fromkeys(warnings):
            print(f"     [check] {w}")
        if opts.print:
            print(json.dumps(clean_res, indent=2, ensure_ascii=False))
        print(f"[ok] {Path(opts.out) / (p.stem + '.json')}  +  .xlsx")


if __name__ == "__main__":
    main()
